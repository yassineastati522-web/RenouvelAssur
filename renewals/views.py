from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required, user_passes_test
from django.db import transaction
from django.db.models import Case, Count, DateField, DateTimeField, F, OuterRef, Prefetch, Q, Subquery, Sum, Value, When
from django.db.models.functions import Coalesce
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.crypto import constant_time_compare
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .audit import record_audit_event
from .forms import (
    ChecklistDateFilterForm,
    ClientForm,
    ExpiredDateFilterForm,
    ImportForm,
    InteractionForm,
    ProvisionalPlanForm,
)
from .models import (
    AuditEvent,
    CallInteraction,
    Client,
    Contract,
    ImportBatch,
    QUICK_CALL_RESULTS,
    Termination,
)
from .services import import_contracts


def scoped_contracts(user):
    qs = Contract.objects.select_related(
        "client",
        "assigned_agent",
        "termination",
    ).prefetch_related(Prefetch(
        "interactions",
        queryset=CallInteraction.objects.order_by("-occurred_at", "-pk")[:1],
        to_attr="_latest_interactions",
    ))
    return qs if user.is_agency_admin else qs.filter(Q(assigned_agent=user) | Q(assigned_agent__isnull=True))


def exclude_terminated_contracts(qs):
    return qs.exclude(
        Q(renewal_status=Contract.RenewalStatus.TERMINATED)
        | Q(manually_terminated=True)
        | Q(termination__isnull=False)
    )


def apply_search(qs, request):
    query = request.GET.get("q", "").strip()
    if query:
        qs = qs.filter(Q(client__name__icontains=query) | Q(client__phone__icontains=query) |
            Q(policy_number__icontains=query) | Q(registration__icontains=query) |
            Q(receipt__icontains=query))
    return qs


def paginate(request, qs, per_page=25):
    return Paginator(qs, per_page).get_page(request.GET.get("page"))


def _excel_safe_text(value):
    text = "" if value is None else str(value)
    if text.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def export_expired_contracts(qs, date_from, date_to):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Non renouvelés"
    headers = [
        "Assuré",
        "Téléphone",
        "Police",
        "Quittance",
        "Marque",
        "Immatriculation",
        "Prime TTC (MAD)",
        "Date d'échéance",
        "Statut",
        "Dernier appel",
        "Date du dernier appel",
        "Agent",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="08765F")

    for contract in qs.order_by("end_date", "client__name", "pk"):
        last_interaction = contract.last_interaction
        agent_name = ""
        if contract.assigned_agent:
            agent_name = (
                contract.assigned_agent.get_full_name()
                or contract.assigned_agent.username
            )
        sheet.append([
            _excel_safe_text(contract.client.name),
            _excel_safe_text(contract.client.phone),
            _excel_safe_text(contract.policy_number),
            _excel_safe_text(contract.receipt),
            _excel_safe_text(contract.brand),
            _excel_safe_text(contract.registration),
            contract.total_premium,
            contract.end_date,
            contract.get_renewal_status_display(),
            (
                last_interaction.get_call_result_display()
                if last_interaction
                else "Pas encore appelé"
            ),
            (
                timezone.localtime(last_interaction.occurred_at).replace(tzinfo=None)
                if last_interaction
                else None
            ),
            _excel_safe_text(agent_name),
        ])

    for column in ("B", "C", "D", "F"):
        for cell in sheet[column][1:]:
            cell.number_format = "@"
    for cell in sheet["G"][1:]:
        cell.number_format = '#,##0.00'
    for cell in sheet["H"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in sheet["K"][1:]:
        cell.number_format = "dd/mm/yyyy hh:mm"
    widths = (28, 18, 20, 18, 16, 20, 18, 18, 24, 24, 22, 20)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="non-renouveles_'
        f'{date_from:%Y-%m-%d}_{date_to:%Y-%m-%d}.xlsx"'
    )
    return response


@login_required
def dashboard(request):
    today = timezone.localdate(); qs = scoped_contracts(request.user)
    soon = exclude_terminated_contracts(
        qs.filter(end_date__range=(today, today + timedelta(days=15)))
    ).exclude(renewal_status=Contract.RenewalStatus.RENEWED)
    expired = exclude_terminated_contracts(
        qs.filter(end_date__lt=today)
    ).exclude(renewal_status=Contract.RenewalStatus.RENEWED)
    due_today = CallInteraction.objects.filter(contract__in=qs, next_follow_up__date__lte=today).values("contract").distinct().count()
    renewed = qs.filter(renewal_status=Contract.RenewalStatus.RENEWED)
    stats = {
        "soon": soon.count(), "due_today": due_today,
        "answered": CallInteraction.objects.filter(contract__in=qs, call_result=CallInteraction.Result.ANSWERED).values("contract").distinct().count(),
        "unreachable": qs.filter(renewal_status=Contract.RenewalStatus.UNREACHABLE).count(), "renewed": renewed.count(),
        "not_renewed": expired.count(), "terminated": qs.filter(renewal_status=Contract.RenewalStatus.TERMINATED).count(),
        "at_risk": soon.aggregate(v=Sum("total_premium"))["v"] or Decimal("0"),
        "renewed_premium": renewed.aggregate(v=Sum("total_premium"))["v"] or Decimal("0"),
    }
    return render(request, "renewals/dashboard.html", {"stats": stats, "upcoming": soon[:8], "followups": qs.filter(interactions__next_follow_up__date__lte=today).distinct()[:6]})


@login_required
def contract_list(request):
    today = timezone.localdate()
    period = request.GET.get("period", "plus15")
    period_ranges = {
        "minus15": (today - timedelta(days=15), today - timedelta(days=1)),
        "minus7": (today - timedelta(days=7), today - timedelta(days=1)),
        "plus7": (today, today + timedelta(days=7)),
        "plus15": (today, today + timedelta(days=15)),
    }
    if period not in {*period_ranges, "all"}:
        period = "plus15"

    selected_status = request.GET.get("status", "not_renewed")
    if selected_status not in {"renewed", "not_renewed"}:
        selected_status = "not_renewed"

    qs = exclude_terminated_contracts(scoped_contracts(request.user))
    if period != "all":
        qs = qs.filter(end_date__range=period_ranges[period])

    if selected_status == "renewed":
        qs = qs.filter(renewal_status=Contract.RenewalStatus.RENEWED)
    else:
        qs = qs.exclude(
            Q(renewal_status=Contract.RenewalStatus.RENEWED)
            | Q(renewed_contract__isnull=False)
        )

    qs = apply_search(qs, request)
    return render(request, "renewals/contract_list.html", {
        "contracts": paginate(request, qs),
        "period": period,
        "selected_status": selected_status,
        "title": "Contrats à renouveler",
    })


@login_required
def expired_list(request):
    qs = exclude_terminated_contracts(
        scoped_contracts(request.user).filter(end_date__lt=timezone.localdate(), renewed_contract__isnull=True)
    ).exclude(renewal_status=Contract.RenewalStatus.RENEWED)
    date_filter_form = ExpiredDateFilterForm(request.GET or None)
    export_requested = request.GET.get("action") == "export"
    if date_filter_form.is_valid():
        date_from = date_filter_form.cleaned_data.get("date_from")
        date_to = date_filter_form.cleaned_data.get("date_to")
        if date_from:
            qs = qs.filter(end_date__gte=date_from)
        if date_to:
            qs = qs.filter(end_date__lte=date_to)
        if export_requested:
            if not date_from or not date_to:
                date_filter_form.add_error(
                    None,
                    "Veuillez sélectionner un intervalle de dates avant l’export.",
                )
            elif (date_to - date_from).days > 6:
                date_filter_form.add_error(
                    None,
                    "L’intervalle d’export ne peut pas dépasser 7 jours.",
                )
            else:
                exported_count = qs.count()
                response = export_expired_contracts(qs, date_from, date_to)
                record_audit_event(
                    actor=request.user,
                    action=AuditEvent.Action.DATA_EXPORTED,
                    details={
                        "export_type": "non_renewed_contracts",
                        "date_from": date_from.isoformat(),
                        "date_to": date_to.isoformat(),
                        "exported_count": exported_count,
                    },
                )
                return response
    return render(request, "renewals/contract_list.html", {
        "contracts": paginate(request, qs),
        "title": "Clients non renouvelés",
        "expired": True,
        "date_filter_form": date_filter_form,
    })


@login_required
def terminated_list(request):
    accessible_contracts = scoped_contracts(request.user)
    qs = Termination.objects.select_related("contract__client").filter(
        contract__in=accessible_contracts,
    ).annotate(
        client_terminated_count=Count(
            "contract__client__contracts__termination",
            filter=Q(contract__client__contracts__in=accessible_contracts),
            distinct=True,
        ),
    ).order_by("-date", "-pk")
    missing_premium_count = qs.filter(premium__isnull=True).count()
    return render(request, "renewals/terminated_list.html", {
        "terminations": paginate(request, qs),
        "missing_premium_count": missing_premium_count,
    })


@login_required
@transaction.atomic
def contract_detail(request, pk):
    contracts = scoped_contracts(request.user)
    if request.method == "POST":
        contracts = contracts.select_for_update(of=("self",))
    contract = get_object_or_404(contracts, pk=pk)
    is_plan_post = (
        request.method == "POST"
        and request.POST.get("form_action") == "provisional_plan"
    )
    if request.method == "POST" and not is_plan_post and not constant_time_compare(
        request.POST.get("call_context", ""), contract.call_context_token,
    ):
        messages.error(request, "Formulaire périmé ou incomplet : actualisez la fiche "
                       "avant de valider l’appel. Aucun appel enregistré.")
        return redirect("contract_detail", pk=contract.pk)
    form = InteractionForm(
        request.POST if request.method == "POST" and not is_plan_post else None,
        initial={"renewal_status": contract.renewal_status},
    )
    provisional_plan_form = ProvisionalPlanForm(
        request.POST if is_plan_post else None,
        contract=contract,
    )
    if is_plan_post and not contract.is_provisional:
        messages.error(request, "Ce contrat n’a pas de suivi provisoire actif.")
        return redirect("contract_detail", pk=contract.pk)
    if is_plan_post and provisional_plan_form.is_valid():
        selected_count = provisional_plan_form.cleaned_data[
            "provisional_selected_count"
        ]
        contract.provisional_selected_count = selected_count
        contract.save(
            update_fields=["provisional_selected_count", "updated_at"]
        )
        record_audit_event(
            actor=request.user,
            action=AuditEvent.Action.PROVISIONAL_PLAN_UPDATED,
            target=contract,
            details={"provisional_count": selected_count},
        )
        messages.success(
            request,
            f"Choix enregistré : {selected_count} provisoire"
            f"{'s' if selected_count > 1 else ''} pour ce contrat.",
        )
        return redirect("contract_detail", pk=contract.pk)
    if request.method == "POST" and not is_plan_post and form.is_valid():
        selected_status = form.cleaned_data["renewal_status"]
        previous_status = contract.renewal_status
        already_terminated = (
            contract.is_terminated
            or Termination.objects.filter(contract=contract).exists()
        )
        if (
            already_terminated
            and selected_status != Contract.RenewalStatus.TERMINATED
        ):
            messages.error(
                request,
                "Un contrat résilié ne peut pas être rouvert depuis la fiche "
                "d’appel. Contactez un administrateur pour corriger ses données.",
            )
            return redirect("contract_detail", pk=contract.pk)
        interaction = form.save(commit=False); interaction.contract = contract; interaction.employee = request.user; interaction.save()
        record_audit_event(
            actor=request.user,
            action=AuditEvent.Action.CALL_RECORDED,
            target=contract,
            details={"outcome": interaction.call_result},
        )
        contract.renewal_status = interaction.renewal_status
        update_fields = ["renewal_status", "updated_at"]
        if interaction.renewal_status == Contract.RenewalStatus.TERMINATED:
            cancellation_date = timezone.localdate()
            contract.end_date = cancellation_date
            update_fields.append("end_date")
            Termination.objects.update_or_create(
                contract=contract,
                defaults={
                    "date": cancellation_date,
                    "reason": interaction.comment.strip()
                    or "Résiliation enregistrée lors d’un appel",
                    "recorded_by": request.user,
                },
            )
        contract.save(update_fields=update_fields)
        record_audit_event(
            actor=request.user,
            action=AuditEvent.Action.CONTRACT_UPDATED,
            target=contract,
            details={
                "previous_status": previous_status,
                "new_status": interaction.renewal_status,
            },
        )
        if interaction.renewal_status == Contract.RenewalStatus.TERMINATED:
            record_audit_event(
                actor=request.user,
                action=AuditEvent.Action.TERMINATION_RECORDED,
                target=contract,
            )
        messages.success(request, "Interaction enregistrée dans l’historique.")
        missed = {CallInteraction.Result.VOICEMAIL, CallInteraction.Result.UNREACHABLE, CallInteraction.Result.OFF}
        missed_days = contract.interactions.filter(call_result__in=missed).dates("occurred_at", "day").count()
        if missed_days >= 3 and contract.renewal_status != Contract.RenewalStatus.UNREACHABLE:
            messages.warning(request, "Trois tentatives sans réponse sur des jours distincts : le statut « Injoignable » est suggéré.")
        return redirect("contract_detail", pk=contract.pk)
    return render(request, "renewals/contract_detail.html", {
        "contract": contract,
        "form": form,
        "provisional_plan_form": provisional_plan_form,
    })


@login_required
@user_passes_test(lambda u: u.is_agency_admin)
@require_http_methods(["GET", "POST"])
@transaction.atomic
def contract_delete(request, pk):
    contract = get_object_or_404(Contract.objects.select_related("client"), pk=pk)
    if request.method == "POST":
        policy_number = contract.policy_number
        client_name = contract.client.name
        record_audit_event(
            actor=request.user,
            action=AuditEvent.Action.CONTRACT_DELETED,
            target=contract,
        )
        contract.delete()
        messages.success(request, f"Le contrat {policy_number} de {client_name} a été supprimé.")
        return redirect("contract_list")
    return render(request, "renewals/contract_confirm_delete.html", {
        "contract": contract,
        "interaction_count": contract.interactions.count(),
        "has_termination": Termination.objects.filter(contract=contract).exists(),
    })


@login_required
@transaction.atomic
def call_checklist(request):
    allowed_contracts = exclude_terminated_contracts(
        scoped_contracts(request.user)
    )
    if request.method == "POST":
        contract = get_object_or_404(
            allowed_contracts.select_for_update(of=("self",)),
            pk=request.POST.get("contract"),
        )
        if not constant_time_compare(
            request.POST.get("call_context", ""), contract.call_context_token,
        ):
            messages.error(request, "Formulaire périmé ou incomplet : actualisez la checklist "
                           "avant de valider l’appel. Aucun appel enregistré.")
            return redirect(request.get_full_path())
        result = request.POST.get("call_result", "")
        allowed_results = {value for value, _label in QUICK_CALL_RESULTS}
        if result not in allowed_results:
            messages.error(request, "Choisissez un résultat d’appel valide.")
        else:
            CallInteraction.objects.create(
                contract=contract,
                employee=request.user,
                channel=CallInteraction.Channel.PHONE,
                call_result=result,
                renewal_status=contract.renewal_status,
                comment=request.POST.get("comment", "").strip(),
            )
            record_audit_event(
                actor=request.user,
                action=AuditEvent.Action.CALL_RECORDED,
                target=contract,
                details={"outcome": result},
            )
            messages.success(request, f"Appel de {contract.client.name} enregistré dans la checklist.")
        return redirect(request.get_full_path())

    now = timezone.now()
    latest_call = CallInteraction.objects.filter(
        contract=OuterRef("pk"),
        channel=CallInteraction.Channel.PHONE,
    ).annotate(
        cycle_start=Case(
            When(
                contract__renewal_calls_started_at__lte=now,
                then=F("contract__renewal_calls_started_at"),
            ),
            When(
                contract__is_provisional=True,
                then=F("contract__provisional_calls_started_at"),
            ),
            default=Value(None),
            output_field=DateTimeField(),
        ),
    ).filter(
        Q(cycle_start__isnull=True) | Q(occurred_at__gte=F("cycle_start"))
    ).order_by("-occurred_at", "-pk")
    # Le compteur et le dernier appel portent sur la même étape de suivi.
    cycle_call_count = latest_call.order_by().values("contract").annotate(
        total=Count("pk"),
    ).values("total")
    closed_statuses = [
        Contract.RenewalStatus.RENEWED,
        Contract.RenewalStatus.TERMINATED,
        Contract.RenewalStatus.REFUSED,
        Contract.RenewalStatus.COMPETITOR,
    ]
    contracts = exclude_terminated_contracts(
        allowed_contracts.exclude(renewal_status__in=closed_statuses)
    ).annotate(
        action_date=Case(
            When(renewal_calls_started_at__lte=now, then=F("end_date")),
            When(
                is_provisional=True,
                provisional_due_date__isnull=False,
                then=F("provisional_due_date"),
            ),
            default=F("end_date"),
            output_field=DateField(),
        ),
        last_call_result=Subquery(latest_call.values("call_result")[:1]),
        last_call_at=Subquery(latest_call.values("occurred_at")[:1]),
        call_attempts=Coalesce(Subquery(cycle_call_count), Value(0)),
    ).order_by("action_date", "client__name", "pk")

    today = timezone.localdate()
    date_filter_form = ChecklistDateFilterForm(request.GET or None)
    selected_due_date = None
    if date_filter_form.is_valid():
        selected_due_date = date_filter_form.cleaned_data.get("due_date")

    due_filter = request.GET.get("due_filter", "all")
    if selected_due_date:
        contracts = contracts.filter(action_date__gte=selected_due_date)
    elif due_filter == "gt7":
        contracts = contracts.filter(action_date__gt=today + timedelta(days=7))
    elif due_filter == "gt15":
        contracts = contracts.filter(action_date__gt=today + timedelta(days=15))
    else:
        due_filter = "all"
        contracts = contracts.filter(action_date__gte=today)

    total_count = contracts.count()
    pending_count = contracts.filter(last_call_at__isnull=True).count()
    unavailable_results = [
        CallInteraction.Result.VOICEMAIL,
        CallInteraction.Result.UNREACHABLE,
    ]
    unavailable_count = contracts.filter(
        last_call_result__in=unavailable_results,
    ).count()
    completed_count = contracts.filter(
        last_call_at__isnull=False,
    ).exclude(last_call_result__in=unavailable_results).count()
    call_status = request.GET.get("call_status", "all")
    if call_status == "pending":
        contracts = contracts.filter(last_call_at__isnull=True)
    elif call_status == "completed":
        contracts = contracts.filter(
            last_call_at__isnull=False,
        ).exclude(last_call_result__in=unavailable_results)
    elif call_status == "unavailable":
        contracts = contracts.filter(last_call_result__in=unavailable_results)
    else:
        call_status = "all"

    page = paginate(request, contracts, per_page=30)
    result_labels = dict(QUICK_CALL_RESULTS)
    for contract in page:
        contract.last_call_label = result_labels.get(contract.last_call_result, "À appeler")
        contract.display_days_remaining = (contract.action_date - today).days

    return render(request, "renewals/call_checklist.html", {
        "contracts": page,
        "call_results": QUICK_CALL_RESULTS,
        "call_status": call_status,
        "due_filter": due_filter,
        "date_filter_form": date_filter_form,
        "total_count": total_count,
        "pending_count": pending_count,
        "completed_count": completed_count,
        "unavailable_count": unavailable_count,
    })


@login_required
def client_list(request):
    if request.user.is_agency_admin:
        qs = Client.objects.annotate(
            contract_count=Count("contracts", distinct=True),
        )
    else:
        accessible_contract_ids = (
            scoped_contracts(request.user).order_by().values("pk")
        )
        qs = Client.objects.annotate(
            contract_count=Count(
                "contracts",
                filter=Q(contracts__in=accessible_contract_ids),
                distinct=True,
            ),
        ).filter(contract_count__gt=0)
    qs = qs.order_by("name", "pk")
    query = request.GET.get("q", "").strip()
    if query: qs = qs.filter(Q(name__icontains=query) | Q(phone__icontains=query) | Q(external_id__icontains=query))
    return render(request, "renewals/client_list.html", {"clients": paginate(request, qs)})


@login_required
@transaction.atomic
def client_detail(request, pk):
    client = get_object_or_404(Client, pk=pk)
    allowed = scoped_contracts(request.user).filter(client=client)
    if not allowed.exists() and not request.user.is_agency_admin: return HttpResponseForbidden()
    form = ClientForm(request.POST or None, instance=client)
    if request.method == "POST" and form.is_valid():
        changed_fields = form.changed_data
        form.save()
        record_audit_event(
            actor=request.user,
            action=AuditEvent.Action.CLIENT_UPDATED,
            target=client,
            details={"changed_fields": changed_fields},
        )
        messages.success(request, "Coordonnées mises à jour.")
        return redirect("client_detail", pk=pk)
    interactions = paginate(
        request,
        CallInteraction.objects.filter(contract__in=allowed)
        .select_related("contract", "employee")
        .order_by("-occurred_at", "-pk"),
        per_page=30,
    )
    return render(request, "renewals/client_detail.html", {"client": client, "contracts": allowed, "interactions": interactions, "form": form})


@login_required
@transaction.atomic
def import_view(request):
    allowed_types = {
        ImportBatch.ImportType.UPCOMING,
        ImportBatch.ImportType.BORDEREAU,
        ImportBatch.ImportType.PROVISIONAL,
    }
    requested_type = request.POST.get("import_kind") if request.method == "POST" else None
    upcoming_form = ImportForm(
        request.POST if requested_type == ImportBatch.ImportType.UPCOMING else None,
        request.FILES if requested_type == ImportBatch.ImportType.UPCOMING else None,
        prefix="upcoming",
        import_type=ImportBatch.ImportType.UPCOMING,
    )
    bordereau_form = ImportForm(
        request.POST if requested_type == ImportBatch.ImportType.BORDEREAU else None,
        request.FILES if requested_type == ImportBatch.ImportType.BORDEREAU else None,
        prefix="bordereau",
        import_type=ImportBatch.ImportType.BORDEREAU,
    )
    provisional_form = ImportForm(
        request.POST if requested_type == ImportBatch.ImportType.PROVISIONAL else None,
        request.FILES if requested_type == ImportBatch.ImportType.PROVISIONAL else None,
        prefix="provisional",
        import_type=ImportBatch.ImportType.PROVISIONAL,
    )

    if request.method == "POST":
        if requested_type not in allowed_types:
            messages.error(request, "Choisissez l’un des trois types d’importation.")
        else:
            active_form = {
                ImportBatch.ImportType.UPCOMING: upcoming_form,
                ImportBatch.ImportType.BORDEREAU: bordereau_form,
                ImportBatch.ImportType.PROVISIONAL: provisional_form,
            }[requested_type]
            if active_form.is_valid():
                upload = active_form.cleaned_data["file"]
                try:
                    batch = import_contracts(
                        upload,
                        request.user,
                        expected_type=requested_type,
                    )
                except ValueError as exc:
                    active_form.add_error("file", str(exc))
                    record_audit_event(
                        actor=request.user,
                        action=AuditEvent.Action.IMPORT_FAILED,
                        details={
                            "import_type": requested_type,
                            "outcome": "validation_error",
                        },
                    )
                else:
                    internal_error = any(
                        error.get("code") == "internal_error"
                        for error in batch.errors
                    )
                    if internal_error:
                        messages.error(
                            request,
                            "Import annulé à cause d’une erreur interne. "
                            "Aucun contrat n’a été modifié.",
                        )
                    else:
                        messages.success(
                            request,
                            f"{batch.get_import_type_display()} importé : "
                            f"{batch.added_rows} ajout(s), "
                            f"{batch.updated_rows} mise(s) à jour, "
                            f"{batch.rejected_rows} rejet(s).",
                        )
                    record_audit_event(
                        actor=request.user,
                        action=(
                            AuditEvent.Action.IMPORT_FAILED
                            if internal_error
                            else AuditEvent.Action.IMPORT_COMPLETED
                        ),
                        target=batch,
                        details={
                            "import_type": batch.import_type,
                            "added_rows": batch.added_rows,
                            "updated_rows": batch.updated_rows,
                            "rejected_rows": batch.rejected_rows,
                            "outcome": (
                                "rolled_back" if internal_error else "completed"
                            ),
                        },
                    )
                    return redirect("import_report", pk=batch.pk)

    return render(request, "renewals/import.html", {
        "upcoming_form": upcoming_form,
        "bordereau_form": bordereau_form,
        "provisional_form": provisional_form,
        "imports": ImportBatch.objects.all().order_by("-imported_at")[:20],
    })


@login_required
def import_report(request, pk):
    return render(request, "renewals/import_report.html", {"batch": get_object_or_404(ImportBatch, pk=pk)})
